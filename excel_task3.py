import os
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import *
from openpyxl.worksheet.table import Table, TableStyleInfo

#Get the directory 
BASE_DIR = os.path.dirname(os.path.realpath(__file__))
#Get the text file
text_file = os.path.join(BASE_DIR,'Employee.txt')

#this list to store the data inside the text file as a list
records = []

#open file and 
with open(text_file) as my_text:
    #make sure that we read the file from the beginning
    my_text.seek(0)
    for record in my_text:
        records.append(record.rstrip("\n").split(';')) 

#now we create our workbook to create the file
workbook = Workbook()

#Path for creating our excel file
excel_file = os.path.join(BASE_DIR,'EmployeeTest.xlsx')


#workbook.save(excel_file)

#Find sheets inside the excel file
print(workbook.sheetnames)


#Reference to the sheet we want to work with
sheet = workbook['Sheet']
#Rename the sheet
sheet.title = "Employees"

#Adding the records to the file
for row in records:
    sheet.append(row)

#Create a table in our sheet Based on our data
table = Table(displayName='Table',ref="A1:G11") #We have a 7 cols and 11 rows

#Define a style for the table and rows,cols stripes
#Find all table style regarding openpyxl
print(openpyxl.worksheet.table.TABLESTYLES)

#define a style for the table
table_style = TableStyleInfo(name="TableStyleMedium23",showRowStripes=True, showColumnStripes=True)

#add the style to the table
table.tableStyleInfo = table_style

#add table to the sheet
sheet.add_table(table)

#showing the highest salaries employees in different color and bold and italic
font = Font(color=colors.RED,bold=True,italic=True)

#add the font and font styles to the sheet
for cell_number in range(2,11):
    if int(sheet[f'G{cell_number}'].value) > 55000  :
        sheet[f'G{cell_number}'].font = font
    
#Last thing to do is saving the changes to our file
workbook.save(excel_file)
    
#Close the workbook
workbook.close()




